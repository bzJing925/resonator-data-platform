"""Read actual complete CSV/XLSX artifacts; large gates never mock row counts."""
import csv
import json
import threading
import zipfile
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import insert, select
from sqlalchemy.orm import Session

from app.models import Batch, Device
from app.schemas.export_job import ExportCreate
from app.services import export_job_service as jobs
from app.services import export_writer as writer
from app.services.query_snapshot import QuerySnapshot
from app.services.source_identity import assign_batch_identity
from app.workers.export_job import execute_export
from tests.fixtures.platform_db import DATABASES, platform_database


@pytest.fixture(params=DATABASES)
def export_database(request, tmp_path, monkeypatch):
    from app.config import Settings
    settings = Settings(DATA_ROOT=tmp_path / 'data', ALN_DESKTOP_MODE=False,
                        WORKER_HEARTBEAT_SECONDS=1)
    for module in ('app.config', 'app.services.export_job_service',
                   'app.services.export_writer', 'app.workers.export_job',
                   'app.services.artifact_lifecycle'):
        monkeypatch.setattr(module + '.get_settings', lambda: settings)
    with platform_database(request.param, tmp_path) as (db, batch):
        assign_batch_identity(db, batch)
        batch.result_revision = 1
        db.commit()
        yield db, batch, settings


def seed(db, batch, count):
    for start in range(0, count, 2000):
        db.execute(insert(Device), [{'batch_id': batch.id, 'original_filename': f'{i}.s1p',
            'source_device_key': str(UUID(int=i + 1)), 'qs': -1.0}
            for i in range(start, min(count, start + 2000))])
        db.commit()  # Fixture setup is not the publication transaction under test.
    batch.device_count = count
    db.commit()


def admit(db, format='csv', fields=None, filters=None):
    request = ExportCreate.model_validate({'request_key': str(uuid4()), 'format': format,
        'query_spec': {'filters': filters or {}, 'fields': fields or ['id'], 'order_by': None},
        'expected_sources': []})
    result, _, token = jobs.create_job(db.bind, request)
    return result['id'], token


def run(db, job_id, token):
    return execute_export(db.bind, job_id, token, 'test-export-writer', writer.write_export)


def test_actual_csv_200001_ids_and_zip(export_database):
    db, batch, settings = export_database
    seed(db, batch, 200001)
    expected = set(db.scalars(select(Device.id)).all())
    db.rollback()
    job_id, token = admit(db)
    assert run(db, job_id, token)['published']
    status = jobs.get_job(db.bind, job_id)
    root = settings.exports_dir / job_id
    with (root / 'data.csv').open(newline='') as file:
        ids = [int(row['id']) for row in csv.DictReader(file)]
    assert len(ids) == 200001 and set(ids) == expected
    manifest = json.loads((root / 'manifest.json').read_text())
    assert manifest['returned'] == len(ids) == status['rows'] == manifest['matched_total']
    assert manifest['total_relation'] == 'exact' and not manifest['truncated']
    assert manifest['sampling']['method'] == 'none'
    assert str(settings.DATA_ROOT) not in json.dumps(manifest)
    with zipfile.ZipFile(root / 'data.zip') as archive:
        assert set(archive.namelist()) == {'data.csv', 'manifest.json'}
        assert archive.read('data.csv') == (root / 'data.csv').read_bytes()
    assert not list(root.rglob('*.part'))


def test_real_xlsx_small_sheet_boundary_formula_and_numbers(export_database, monkeypatch):
    db, batch, settings = export_database
    seed(db, batch, 4)
    first = db.scalar(select(Device).order_by(Device.id))
    first.display_name = '=HYPERLINK("evil")'
    db.commit()
    monkeypatch.setattr(writer, 'XLSX_SHEET_ROWS', 3)
    job_id, token = admit(db, 'xlsx', fields=['id', 'display_name', 'qs'])
    assert run(db, job_id, token)['published']
    root = settings.exports_dir / job_id
    workbook = load_workbook(root / 'data.xlsx', read_only=True, data_only=False)
    assert workbook.sheetnames == ['说明', 'data_1', 'data_2']
    a, b = list(workbook['data_1'].values), list(workbook['data_2'].values)
    assert len(a) == 4 and len(b) == 2
    assert a[1][1].startswith("'=") and a[1][2] == -1 and isinstance(a[1][2], int | float)
    assert jobs.get_job(db.bind, job_id)['manifest']['formula_escaped_cells'] == 1
    workbook.close()


@pytest.mark.parametrize('failure', ['source', 'new_batch', 'cancel', 'io', 'rows', 'bytes'])
def test_failure_never_publishes_partial(export_database, monkeypatch, failure):
    db, batch, settings = export_database
    seed(db, batch, 3)
    job_id, token = admit(db)
    original = QuerySnapshot.partitions
    def changed(snapshot):
        chunks = original(snapshot)
        for chunk in chunks:
            if failure in {'source', 'new_batch'}:
                with Session(db.bind) as change:
                    if failure == 'source':
                        change.scalar(select(Device).order_by(Device.id)).qs = 9
                    else:
                        added = Batch(batch_no='new-match', mapping_id=batch.mapping_id,
                                      file_path='files')
                        change.add(added)
                        change.flush()
                        assign_batch_identity(change, added)
                        change.add(Device(batch_id=added.id, original_filename='new.s1p',
                                          source_device_key=str(uuid4())))
                    change.commit()
            elif failure == 'cancel':
                jobs.request_cancel(db.bind, job_id)
            elif failure == 'io':
                raise OSError('injected disk error')
            yield chunk
    if failure == 'rows':
        settings.EXPORT_MAX_ROWS = 2
    elif failure == 'bytes':
        settings.EXPORT_MAX_BYTES = 8
    else:
        monkeypatch.setattr(QuerySnapshot, 'partitions', changed)
    result = run(db, job_id, token)
    assert not result['published']
    status = jobs.get_job(db.bind, job_id)
    assert status['status'] == ('cancelled' if failure == 'cancel' else 'failed')
    assert status['artifacts'] == [] and status['manifest'] is None
    assert not list((settings.exports_dir / job_id).rglob('*.part'))
    assert not (settings.exports_dir / job_id / 'data.csv').exists()
    if failure in {'source', 'new_batch'}:
        assert status['error']['code'] == 'SOURCE_CHANGED'


def test_change_after_writer_before_terminal_commit(export_database):
    db, batch, settings = export_database
    seed(db, batch, 2)
    job_id, token = admit(db)
    ticket = jobs.claim_job(db.bind, job_id, token, 'writer')
    manifest, artifacts = writer.write_export(db.bind, ticket, token, threading.Event())
    with Session(db.bind) as change:
        change.scalar(select(Device).order_by(Device.id)).qs = 77
        change.commit()
    from app.services.query_snapshot import ExportWriteError
    with pytest.raises(ExportWriteError, match='SOURCE_CHANGED'):
        jobs.finish_job(db.bind, job_id, token, 'success', manifest=manifest, artifacts=artifacts)
    assert jobs.get_job(db.bind, job_id)['status'] == 'running'
    assert jobs.get_job(db.bind, job_id)['artifacts'] == []


@pytest.mark.integration
def test_actual_excel_capacity_boundary_sqlite(tmp_path, monkeypatch):
    from app.config import Settings
    settings = Settings(DATA_ROOT=tmp_path / 'data', WORKER_HEARTBEAT_SECONDS=1,
                        ALN_DESKTOP_MODE=False, EXPORT_LEASE_SECONDS=180)
    for module in ('app.config', 'app.services.export_job_service',
                   'app.services.export_writer', 'app.workers.export_job',
                   'app.services.artifact_lifecycle'):
        monkeypatch.setattr(module + '.get_settings', lambda: settings)
    with platform_database('sqlite', tmp_path) as (db, batch):
        assign_batch_identity(db, batch)
        batch.result_revision = 1
        db.commit()
        seed(db, batch, 1_048_576)
        expected = set(db.scalars(select(Device.id)))
        db.rollback()
        job_id, token = admit(db, 'xlsx')
        assert run(db, job_id, token)['published']
        workbook = load_workbook(settings.exports_dir / job_id / 'data.xlsx', read_only=True)
        assert workbook.sheetnames == ['说明', 'data_1', 'data_2']
        ids, counts = set(), []
        for sheet_name in ['data_1', 'data_2']:
            count = 0
            for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
                ids.add(row[0])
                count += 1
            counts.append(count)
        workbook.close()
        assert counts == [1_048_575, 1]
        assert ids == expected
        assert jobs.get_job(db.bind, job_id)['manifest']['returned'] == 1_048_576


def test_csv_formula_safety_and_xlsx_byte_budget(export_database):
    db, batch, settings = export_database
    seed(db, batch, 1)
    db.scalar(select(Device)).display_name = ' \t@danger'
    db.commit()
    job_id, token = admit(db, fields=['display_name', 'qs'])
    assert run(db, job_id, token)['published']
    with (settings.exports_dir / job_id / 'data.csv').open(newline='') as file:
        row = next(csv.DictReader(file))
    assert row['display_name'].startswith("'") and float(row['qs']) == -1
    settings.EXPORT_MAX_BYTES = 50
    job_id, token = admit(db, 'xlsx')
    assert not run(db, job_id, token)['published']
    assert jobs.get_job(db.bind, job_id)['status'] == 'failed'
    assert not list((settings.exports_dir / job_id).rglob('*.part'))


def test_publication_locks_exact_origin_batch_pairs(export_database):
    db, batch, _settings = export_database
    seed(db, batch, 1)
    other = Batch(batch_no='other-origin', mapping_id=batch.mapping_id, file_path='files',
                  result_revision=1, origin_id=str(uuid4()), source_batch_id=str(uuid4()))
    db.add(other)
    db.flush()
    db.add(Device(batch_id=other.id, original_filename='other.s1p',
                  source_device_key=str(uuid4())))
    db.add(Batch(batch_no='unmatched-pair', mapping_id=batch.mapping_id, file_path='files',
                 origin_id=batch.origin_id, source_batch_id=other.source_batch_id))
    db.commit()
    job_id, token = admit(db)
    assert run(db, job_id, token)['published']
    assert jobs.get_job(db.bind, job_id)['rows'] == 2
