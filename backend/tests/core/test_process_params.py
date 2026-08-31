"""工艺参数 Excel 解析器测试（合成 workbook，覆盖规格 §5 全部规则）。"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.core.process_params import parse_process_workbook


def _build_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()

    # --- SPM2 ICP Etch：2 个表头块 + mean/spec 行 + 合并单元格 + 百分比格式 ---
    ws = wb.active
    ws.title = "SPM2 ICP Etch"
    ws.append([])  # r1 空行
    ws.append([None, "Information", None, "Measrements"])  # r2
    header = [None, "No.", "Etch Time", "Etch Rate", "Thickness Mapping", None, None, "Comments"]
    ws.append(header)  # r3
    ws.append([None, None, "sec", "nm/min", "THK/A", "WiW 1sig/%", "WtW 1sig/%"])  # r4 单位行
    ws.append([None, "F0425LW1", 98, 24.2, 395.3, 0.029, None, "5000A oxide wafer"])  # r5
    ws["F5"].number_format = "0.00%"
    ws.append([None, "F0425LW2", 98, 24.5, 399.5, 0.031])  # r6
    ws["F6"].number_format = "0.00%"
    ws.append([None, None, "mean", 24.35, 397.4, 0.03])  # r7 mean 汇总行
    ws.append([None, None, None, None, None, "spec. <5%", "spec. <2%"])  # r8 spec 行
    ws.append([])  # r9 空行
    ws.append([None, "No.", "Etch Time"])  # r10 第二表头块
    ws.append([None, None, "sec"])  # r11 单位行
    ws.append([None, "0602A17", 43, 1128.8, 809.0, 0.0351])  # r12
    ws["F12"].number_format = "0.00%"
    ws.merge_cells("G5:G6")  # WtW 组级合并单元格
    ws["G5"] = 0.011
    ws["G5"].number_format = "0.00%"

    # --- SPM3 Mo：group_label + '??' 脏单元格 ---
    ws = wb.create_sheet("SPM3 Mo")
    ws.append([None, "No.", "Ar Flow Rate", "DC Power", "Dep. Time", "Dep. Rate"])  # r1
    ws.append([None, None, "sccm", "kW", "sec", "nm/min"])  # r2 单位行
    ws.append(["DOE2", "F0507A11", 10, 7, 20, 392.3])  # r3
    ws.append([None, "F0426A6", 10, 7, 20, "??"])  # r4 脏单元格

    # --- SPM4 AlScN(30AT)：实验名 run_no ---
    ws = wb.create_sheet("SPM4 AlScN(30AT)")
    ws.append([None, "No.", "N2 Flow Rate", "Ar Flow Rate", "DC Power", "Dep. Time", "Dep. Rate"])
    ws.append([None, None, "sccm", "sccm", "kW", "sec", "nm/min"])
    ws.append(["S21", "DOE1", 40, 20, 7, 307.27, 1.13])

    # --- AlSc：recipe_id + 规格目标行（run_no 为空）+ 中间空行 ---
    ws = wb.create_sheet("AlSc")
    ws.append(["Recipe ID", "No.", "N2 Flow Rate"])
    ws.append([None, None, "sccm"])
    ws.append([None, None, None, None, None, None, None, 3500, "<0.5%"])  # 规格行：B 空、H/I 有内容
    ws.append(["DOE26-3A-180", "0623A07", 46.9])
    ws.append([])  # 空行：静默忽略，不计入 skipped
    ws.append(["DOE35", "0526A22", 46.9])

    wb.save(path)


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "process.xlsx"
    _build_workbook(path)
    return path


def test_icp_rows_and_block_index(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    assert [r["run_no"] for r in data.icp_etch] == ["F0425LW1", "F0425LW2", "0602A17"]
    assert [r["block_index"] for r in data.icp_etch] == [1, 1, 2]
    assert data.icp_etch[0]["comments"] == "5000A oxide wafer"
    assert data.icp_etch[0]["etch_time_sec"] == 98.0


def test_percent_format_multiplied_by_100(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    assert data.icp_etch[0]["wiw_1sig_pct"] == pytest.approx(2.9)
    assert data.icp_etch[2]["wiw_1sig_pct"] == pytest.approx(3.51)


def test_merged_cell_kept_on_anchor_row_only(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    assert data.icp_etch[0]["wtw_1sig_pct"] == pytest.approx(1.1)
    assert data.icp_etch[1]["wtw_1sig_pct"] is None


def test_icp_summary_and_header_rows_skipped(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    rep = data.report.sheets["SPM2 ICP Etch"]
    assert rep.imported == 3
    # r2 Information、r3/r10 表头、r4/r11 单位行、r7 mean、r8 spec → 7 行
    assert rep.skipped == 7


def test_dirty_cell_becomes_none_and_reported(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    assert data.mo_dep[1]["dep_rate_nm_min"] is None
    dirty = [d for d in data.report.dirty_cells if d.sheet == "SPM3 Mo"]
    assert len(dirty) == 1
    assert dirty[0].raw == "??"
    assert dirty[0].cell == "F4"


def test_group_label_and_recipe_id(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    assert data.mo_dep[0]["group_label"] == "DOE2"
    assert data.mo_dep[1]["group_label"] is None
    assert data.aln_dep[0]["group_label"] == "S21"
    assert data.aln_dep[0]["source_sheet"] == "SPM4 AlScN(30AT)"
    alsc = [r for r in data.aln_dep if r["source_sheet"] == "AlSc"]
    assert [r["recipe_id"] for r in alsc] == ["DOE26-3A-180", "DOE35"]
    assert alsc[0]["group_label"] is None


def test_alsc_spec_row_skipped_blank_row_ignored(workbook_path: Path):
    data = parse_process_workbook(workbook_path)
    rep = data.report.sheets["AlSc"]
    assert rep.imported == 2
    # r1 表头、r2 单位行、r3 规格行 → 3 行；r5 空行不计
    assert rep.skipped == 3


def test_missing_sheet_raises(tmp_path: Path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError, match="缺少 sheet"):
        parse_process_workbook(tmp_path / "empty.xlsx")
