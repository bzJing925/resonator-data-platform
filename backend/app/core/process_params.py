"""工艺参数 Excel 解析器（纯函数）。

源文件结构（规格 docs/superpowers/specs/2026-07-20-process-params-design.md §3）：
  SPM2 ICP Etch    ICP 刻蚀（3 个重复表头块，块间有 mean/spec 汇总行）
  SPM3 Mo          Mo 生长（A 列为实验组标签）
  SPM4 AlScN(30AT) AlScN 生长（A 列为实验组标签）
  AlSc             AlSc 生长（A 列为 Recipe ID，含规格目标行）

数据行识别：run_no 列非空且非表头字面量（'No.' / 'Information'）。
合并单元格只返回锚点值（openpyxl 默认行为），组级统计不向下填充。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

_HEADER_TOKENS = {"No.", "Information"}

TEXT_FIELDS = frozenset({"run_no", "group_label", "recipe_id", "comments", "source_sheet"})


@dataclass(frozen=True)
class SheetConfig:
    """一个 sheet 的显式列配置。"""

    sheet: str
    table: str  # 'icp_etch' | 'mo_dep' | 'aln_dep'
    columns: dict[str, str]  # 列字母 -> 字段名
    pct_fields: frozenset[str]  # number_format 含 % 时 ×100 的字段
    has_blocks: bool = False  # 是否有重复表头块（block_index 递增）


ICP_COLUMNS: dict[str, str] = {
    "B": "run_no",
    "C": "etch_time_sec",
    "D": "etch_rate_nm_min",
    "E": "thk_a",
    "F": "wiw_1sig_pct",
    "G": "wtw_1sig_pct",
    "H": "comments",
}

MO_COLUMNS: dict[str, str] = {
    "A": "group_label",
    "B": "run_no",
    "C": "ar_flow_sccm",
    "D": "dc_power_kw",
    "E": "dep_time_sec",
    "F": "dep_rate_nm_min",
    "G": "thk_nm",
    "H": "wiw_1sig_pct",
    "I": "wtw_1sig_pct",
    "J": "stress_avg_mpa",
    "K": "stress_wiw_range_mpa",
    "L": "stress_wtw_range_mpa",
    "M": "rs_ohm_sq",
    "N": "resistivity_uohm_cm",
    "O": "rs_uni_pct",
    "P": "roughness_ra_nm",
    "Q": "xrd_rc_fwhm_deg",
    "R": "comments",
}

# SPM4 / AlSc 共用的列（A 列除外：SPM4 是 group_label，AlSc 是 recipe_id）
ALN_COLUMNS: dict[str, str] = {
    "B": "run_no",
    "C": "n2_flow_sccm",
    "D": "ar_flow_sccm",
    "E": "dc_power_kw",
    "F": "dep_time_sec",
    "G": "dep_rate_nm_min",
    "H": "thk_nm",
    "I": "wiw_1sig_pct",
    "J": "wtw_1sig_pct",
    "K": "stress_avg_mpa",
    "L": "stress_wiw_range_mpa",
    "M": "stress_wtw_range_mpa",
    "N": "refractive_index",
    "O": "roughness_ra_nm",
    "P": "aog_per_um2",
    "Q": "xrd_rc_fwhm_deg",
    "R": "comments",
}

_PCT_FIELDS = frozenset({"wiw_1sig_pct", "wtw_1sig_pct", "rs_uni_pct"})

SHEET_CONFIGS: tuple[SheetConfig, ...] = (
    SheetConfig("SPM2 ICP Etch", "icp_etch", ICP_COLUMNS, _PCT_FIELDS, has_blocks=True),
    SheetConfig("SPM3 Mo", "mo_dep", MO_COLUMNS, _PCT_FIELDS),
    SheetConfig("SPM4 AlScN(30AT)", "aln_dep", {"A": "group_label", **ALN_COLUMNS}, _PCT_FIELDS),
    SheetConfig("AlSc", "aln_dep", {"A": "recipe_id", **ALN_COLUMNS}, _PCT_FIELDS),
)


def _build_table_fields() -> dict[str, frozenset[str]]:
    """每张表的字段全集（同表多 sheet 配置的列并集 + 附加列），缺失字段补 None。"""
    fields: dict[str, set[str]] = {}
    for cfg in SHEET_CONFIGS:
        fields.setdefault(cfg.table, set()).update(cfg.columns.values())
    fields["icp_etch"].add("block_index")
    fields["aln_dep"].add("source_sheet")
    return {k: frozenset(v) for k, v in fields.items()}


TABLE_FIELDS = _build_table_fields()


@dataclass(frozen=True)
class DirtyCell:
    """无法转换为数值而被清洗为 NULL 的单元格。"""

    sheet: str
    cell: str  # 如 'F12'
    raw: str


@dataclass
class SheetReport:
    imported: int = 0
    skipped: int = 0


@dataclass
class ParseReport:
    sheets: dict[str, SheetReport] = field(default_factory=dict)
    dirty_cells: list[DirtyCell] = field(default_factory=list)


@dataclass
class ProcessData:
    icp_etch: list[dict[str, Any]]
    mo_dep: list[dict[str, Any]]
    aln_dep: list[dict[str, Any]]
    report: ParseReport


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (ValueError, TypeError):
        return None


def _is_percent_format(number_format: str | None) -> bool:
    return bool(number_format) and "%" in number_format


def parse_process_workbook(path: str | Path) -> ProcessData:
    """解析工艺参数 xlsx，返回三张表的行数据 + 解析报告。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        missing = [c.sheet for c in SHEET_CONFIGS if c.sheet not in wb.sheetnames]
        if missing:
            raise ValueError(f"Excel 缺少 sheet: {', '.join(missing)}")

        data = ProcessData(icp_etch=[], mo_dep=[], aln_dep=[], report=ParseReport())

        for cfg in SHEET_CONFIGS:
            ws = wb[cfg.sheet]
            rep = SheetReport()
            data.report.sheets[cfg.sheet] = rep
            block_index = 0
            run_no_col = next(col for col, f in cfg.columns.items() if f == "run_no")

            for row in ws.iter_rows():
                # MergedCell 无 column_letter，统一用 get_column_letter
                cells = {get_column_letter(c.column): c for c in row}
                # 整行空白（配置列全无值）→ 静默忽略
                if all(cells.get(col) is None or cells[col].value is None for col in cfg.columns):
                    continue

                run_no_cell = cells.get(run_no_col)
                raw_run_no = run_no_cell.value if run_no_cell is not None else None
                if raw_run_no is None or str(raw_run_no).strip() == "":
                    # mean/spec/单位/规格目标等附属行
                    rep.skipped += 1
                    continue
                run_no_text = str(raw_run_no).strip()
                if run_no_text in _HEADER_TOKENS:
                    if run_no_text == "No." and cfg.has_blocks:
                        block_index += 1
                    rep.skipped += 1
                    continue

                record: dict[str, Any] = {}
                for col, field_name in cfg.columns.items():
                    cell = cells.get(col)
                    value = cell.value if cell is not None else None
                    if field_name in TEXT_FIELDS:
                        text = None if value is None else str(value).strip()
                        record[field_name] = text or None
                        continue
                    num = _to_float(value)
                    if num is None and value is not None and str(value).strip() != "":
                        data.report.dirty_cells.append(
                            DirtyCell(cfg.sheet, f"{col}{row[0].row}", str(value))
                        )
                    if (
                        num is not None
                        and field_name in cfg.pct_fields
                        and _is_percent_format(cell.number_format if cell is not None else None)
                    ):
                        num *= 100
                    record[field_name] = num

                record["run_no"] = run_no_text
                if cfg.table == "aln_dep":
                    record["source_sheet"] = cfg.sheet
                if cfg.has_blocks:
                    record["block_index"] = block_index
                # 同表字段全集对齐（如 AlSc 无 group_label）→ 补 None
                for f in TABLE_FIELDS[cfg.table]:
                    record.setdefault(f, None)
                getattr(data, cfg.table).append(record)
                rep.imported += 1

        return data
    finally:
        wb.close()
