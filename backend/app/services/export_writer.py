"""Bounded streaming exports; only complete files can become public artifacts."""
import csv
import hashlib
import json
import math
import shutil
import threading
import zipfile
from contextlib import contextmanager
from pathlib import PurePosixPath, PureWindowsPath

from sqlalchemy.orm import Session

from app.config import get_settings
from app.services import export_job_service as jobs
from app.services.artifact_lifecycle import acquire_lease, release_lease, renew_lease
from app.services.query_manifest import make_manifest, scope_version
from app.services.query_snapshot import ExportWriteError, export_snapshot
from app.services.query_spec import query_spec

XLSX_SHEET_ROWS = 1_048_575


class BoundedFile:
    """ZIP/XLSX seeking writes count actual on-disk bytes rather than rows or estimates."""
    def __init__(self, path, maximum, stopped=None):
        self.file = path.open('xb')
        self.maximum = maximum
        self.high_water = 0
        self.stopped = stopped

    def write(self, data):
        if self.stopped is not None and self.stopped.is_set():
            raise ExportWriteError('EXPORT_STOPPED')
        if self.file.tell() + len(data) > self.maximum:
            raise ExportWriteError('EXPORT_BYTE_LIMIT')
        count = self.file.write(data)
        self.high_water = max(self.high_water, self.file.tell())
        return count

    def __getattr__(self, name):
        return getattr(self.file, name)

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.file.close()


def safe_cell(value, counters):
    if isinstance(value, float) and not math.isfinite(value):
        counters['nonfinite_values'] += 1
        return None
    if isinstance(value, str) and value.lstrip(' \t\r\n').startswith(('=', '+', '-', '@')):
        counters['formula_escaped_cells'] += 1
        return "'" + value
    return value


def _safe_manifest(value):
    if isinstance(value, dict):
        for item in value.values():
            _safe_manifest(item)
    elif isinstance(value, list):
        for item in value:
            _safe_manifest(item)
    elif isinstance(value, str) and (PurePosixPath(value).is_absolute()
                                   or PureWindowsPath(value).is_absolute()):
        raise ExportWriteError('EXPORT_MANIFEST_PATH_INVALID')


def _check(engine, ticket, token, lost, rows, total):
    if lost.is_set() or not jobs.heartbeat(engine, ticket['id'], token, rows=rows,
                                         progress=int(rows * 99 / max(1, total))):
        raise ExportWriteError('EXPORT_STOPPED')


@contextmanager
def _source_leases(engine, batch_ids, owner, lost):
    leases = []
    stop = threading.Event()
    thread = None
    try:
        for batch_id in batch_ids:
            leases.append(acquire_lease(engine, batch_id, '*', 'export', owner))
        def beat():
            while not stop.wait(max(1, get_settings().ARTIFACT_LEASE_SECONDS / 3)):
                try:
                    if not all(renew_lease(engine, lease, owner) for lease in leases):
                        lost.set()
                        return
                except Exception:
                    lost.set()
                    return
        thread = threading.Thread(target=beat, daemon=True, name='aln-export-source-leases')
        thread.start()
        yield
    finally:
        stop.set()
        if thread:
            thread.join(timeout=1)
        for lease in leases:
            release_lease(engine, lease, owner)


def _csv(snapshot, path, engine, ticket, token, lost, counters, budget):
    rows = 0
    with path.open('x', encoding='utf-8', newline='') as file:
        writer = csv.writer(file)
        writer.writerow(snapshot.fields)
        file.flush()
        if file.tell() > budget:
            raise ExportWriteError('EXPORT_BYTE_LIMIT')
        for chunk in snapshot.partitions():
            _check(engine, ticket, token, lost, rows, snapshot.total)
            for row in chunk:
                writer.writerow([safe_cell(row[field], counters) for field in snapshot.fields])
                if file.tell() > budget:
                    raise ExportWriteError('EXPORT_BYTE_LIMIT')
                rows += 1
        _check(engine, ticket, token, lost, rows, snapshot.total)
    return rows


def _xlsx(snapshot, path, engine, ticket, token, lost, counters, budget):
    from openpyxl import Workbook
    from openpyxl.writer.excel import ExcelWriter
    workbook = Workbook(write_only=True)
    info = workbook.create_sheet('说明')
    info.append(['范围', '完整筛选结果，无图表limit或采样'])
    info.append(['匹配行数', snapshot.total])
    info.append(['epoch', snapshot.epoch])
    info.append(['公式防护', '危险字符串前加单引号；数值保持数值'])
    rows, in_sheet, number, sheet = 0, 0, 0, None
    try:
        for chunk in snapshot.partitions():
            _check(engine, ticket, token, lost, rows, snapshot.total)
            for row in chunk:
                if sheet is None or in_sheet == XLSX_SHEET_ROWS:
                    number += 1
                    sheet = workbook.create_sheet(f'data_{number}')
                    sheet.append(snapshot.fields)
                    in_sheet = 0
                sheet.append([safe_cell(row[field], counters) for field in snapshot.fields])
                in_sheet += 1
                rows += 1
        if sheet is None:
            workbook.create_sheet('data_1').append(snapshot.fields)
        info.append(['实际数据行数', rows])
        _check(engine, ticket, token, lost, rows, snapshot.total)
        with BoundedFile(path, budget, lost) as file:
            with zipfile.ZipFile(file, 'w', compression=zipfile.ZIP_DEFLATED,
                                 allowZip64=True) as archive:
                ExcelWriter(workbook, archive).write_data()
        return rows
    finally:
        # Write-only sheets spool XML in the OS temporary directory; remove it even on failure.
        for worksheet in workbook.worksheets:
            if not worksheet.closed:
                worksheet.close()
            writer = getattr(worksheet, '_writer', None)
            if writer is not None:
                from pathlib import Path
                if Path(writer.out).exists():
                    writer.cleanup()
        workbook.close()


def write_export(engine, ticket, token, lost):
    settings = get_settings()
    root = settings.exports_dir / ticket['id']
    stage = root / f'.{token}'
    stage.mkdir(parents=True, exist_ok=False)
    counters = {'formula_escaped_cells': 0, 'nonfinite_values': 0}
    completed = []
    try:
        with export_snapshot(engine, ticket) as snapshot:
            if snapshot.total > settings.EXPORT_MAX_ROWS:
                raise ExportWriteError('EXPORT_ROW_LIMIT')
            with _source_leases(engine, snapshot.sources.keys(), ticket['id'], lost):
                name = 'data.csv' if ticket['format'] == 'csv' else 'data.xlsx'
                data = stage / (name + '.part')
                writer = _csv if ticket['format'] == 'csv' else _xlsx
                rows = writer(snapshot, data, engine, ticket, token, lost, counters,
                              settings.EXPORT_MAX_BYTES)
                if rows != snapshot.total:
                    raise ExportWriteError('EXPORT_ROW_COUNT_CHANGED')
                manifest = make_manifest(query_spec(**ticket['query_spec']), snapshot.sources,
                    snapshot.verified, rows, rows).model_dump(mode='json')
        with Session(engine) as db:
            version, preference_digest = scope_version(db)
            if (version is None or version[1] != ticket['source_epoch']
                    or preference_digest != ticket['scope_digest']):
                raise ExportWriteError('SOURCE_CHANGED')
        manifest.update(counters)
        manifest['format'] = ticket['format']
        manifest['formula_escape_policy'] = 'prefix_apostrophe_for_string_cells'
        _safe_manifest(manifest)
        encoded = json.dumps(manifest, ensure_ascii=False, allow_nan=False).encode('utf-8')
        metadata = stage / 'manifest.json.part'
        remaining = settings.EXPORT_MAX_BYTES - data.stat().st_size
        if len(encoded) > remaining:
            raise ExportWriteError('EXPORT_BYTE_LIMIT')
        metadata.write_bytes(encoded)
        remaining -= len(encoded)
        parts = [(data, name), (metadata, 'manifest.json')]
        if ticket['format'] == 'csv':
            archive = stage / 'data.zip.part'
            with BoundedFile(archive, remaining, lost) as file:
                with zipfile.ZipFile(file, 'w', compression=zipfile.ZIP_DEFLATED) as bundle:
                    bundle.write(data, 'data.csv')
                    bundle.write(metadata, 'manifest.json')
            parts.append((archive, 'data.zip'))
        _check(engine, ticket, token, lost, rows, rows)
        artifacts = []
        for part, filename in parts:
            final = root / filename
            if final.exists():
                raise ExportWriteError('EXPORT_ARTIFACT_EXISTS')
            part.rename(final)
            completed.append(final)
            with final.open('rb') as file:
                digest = hashlib.file_digest(file, 'sha256').hexdigest()
            artifacts.append({'name': filename, 'bytes': final.stat().st_size, 'sha256': digest})
        return manifest, artifacts
    except Exception:
        for path in completed:
            path.unlink(missing_ok=True)
        raise
    finally:
        shutil.rmtree(stage, ignore_errors=True)
